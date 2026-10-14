"""Tests that source Excel rows are parsed without intermediate CSV conversion."""
import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import Workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
import analyze

class TestExcelInput(unittest.TestCase):
    def test_excel_rows_and_metrics(self):
        with TemporaryDirectory() as folder:
            source = Path(folder) / "sample.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(analyze.REQUIRED)
            ws.append(["101","S1","Example",2,"2011-02-01",10,"0001","United Kingdom"])
            wb.save(source)
            # Exercise the exact reader used in analyze.main
            from openpyxl import load_workbook
            import pandas as pd
            book = load_workbook(source, read_only=True, data_only=True)
            try:
                rows = book.active.values
                columns = next(rows)
                frame = pd.DataFrame(rows, columns=columns)
            finally:
                book.close()
            *_, summary = analyze.analyze(frame)
            self.assertEqual(summary["source_rows"], 1)
            self.assertEqual(summary["gross_sales"], 20.0)

if __name__ == "__main__":
    unittest.main()
